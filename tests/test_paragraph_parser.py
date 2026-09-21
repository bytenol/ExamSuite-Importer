from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from examsuite_importer.parser.paragraph_parser import (
    ParagraphParser,
)


def create_document(tmp_path: Path):
    document = Document()
    paragraph = document.add_paragraph()

    return document, paragraph


def add_text(
    paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    underline: bool = False,
    superscript: bool = False,
    subscript: bool = False,
):
    run = paragraph.add_run(text)

    run.bold = bold
    run.italic = italic
    run.underline = underline

    if superscript:
        run.font.superscript = True

    if subscript:
        run.font.subscript = True

    return run


def test_plain_text():
    document = Document()
    paragraph = document.add_paragraph()

    paragraph.add_run("Hello World")

    result = ParagraphParser(paragraph).parse()

    assert result == "Hello World"


def test_bold_text():
    document = Document()
    paragraph = document.add_paragraph()

    add_text(
        paragraph,
        "bold",
        bold=True,
    )

    result = ParagraphParser(paragraph).parse()

    assert result == "<strong>bold</strong>"


def test_italic_text():
    document = Document()
    paragraph = document.add_paragraph()

    add_text(
        paragraph,
        "italic",
        italic=True,
    )

    result = ParagraphParser(paragraph).parse()

    assert result == "<em>italic</em>"


def test_underlined_text():
    document = Document()
    paragraph = document.add_paragraph()

    add_text(
        paragraph,
        "underline",
        underline=True,
    )

    result = ParagraphParser(paragraph).parse()

    assert result == "<u>underline</u>"


def test_superscript_text():
    document = Document()
    paragraph = document.add_paragraph()

    add_text(
        paragraph,
        "2",
        superscript=True,
    )

    result = ParagraphParser(paragraph).parse()

    assert result == "<sup>2</sup>"


def test_subscript_text():
    document = Document()
    paragraph = document.add_paragraph()

    add_text(
        paragraph,
        "2",
        subscript=True,
    )

    result = ParagraphParser(paragraph).parse()

    assert result == "<sub>2</sub>"


def test_mixed_formatting():
    document = Document()
    paragraph = document.add_paragraph()

    paragraph.add_run("This is ")

    add_text(
        paragraph,
        "bold",
        bold=True,
    )

    paragraph.add_run(" and ")

    add_text(
        paragraph,
        "italic",
        italic=True,
    )

    paragraph.add_run(" text.")

    result = ParagraphParser(paragraph).parse()

    assert result == (
        "This is "
        "<strong>bold</strong>"
        " and "
        "<em>italic</em>"
        " text."
    )


def test_multiple_formatting_on_same_run():
    document = Document()
    paragraph = document.add_paragraph()

    add_text(
        paragraph,
        "formatted",
        bold=True,
        italic=True,
        underline=True,
    )

    result = ParagraphParser(paragraph).parse()

    assert result == (
        "<u><em><strong>formatted</strong></em></u>"
    )


def test_superscript_and_subscript():
    document = Document()
    paragraph = document.add_paragraph()

    paragraph.add_run("H")

    add_text(
        paragraph,
        "2",
        subscript=True,
    )

    paragraph.add_run("O")

    add_text(
        paragraph,
        "2",
        superscript=True,
    )

    result = ParagraphParser(paragraph).parse()

    assert result == (
        "H"
        "<sub>2</sub>"
        "O"
        "<sup>2</sup>"
    )


def test_formatting_is_preserved_in_order():
    document = Document()
    paragraph = document.add_paragraph()

    add_text(
        paragraph,
        "A",
        bold=True,
    )

    add_text(
        paragraph,
        "B",
        italic=True,
    )

    add_text(
        paragraph,
        "C",
        underline=True,
    )

    add_text(
        paragraph,
        "D",
        superscript=True,
    )

    add_text(
        paragraph,
        "E",
        subscript=True,
    )

    result = ParagraphParser(paragraph).parse()

    assert result == (
        "<strong>A</strong>"
        "<em>B</em>"
        "<u>C</u>"
        "<sup>D</sup>"
        "<sub>E</sub>"
    )


def add_equation(paragraph):
    omath = OxmlElement("m:oMath")

    run = OxmlElement("m:r")

    text = OxmlElement("m:t")
    text.text = "x"

    run.append(text)
    omath.append(run)

    paragraph._p.append(omath)

    return omath


def test_equation():
    document = Document()
    paragraph = document.add_paragraph()

    paragraph.add_run("Solve ")

    add_equation(paragraph)

    paragraph.add_run(" = 2")

    result = ParagraphParser(paragraph).parse()

    assert "x" in result
    assert "Solve" in result
    assert "= 2" in result


def test_empty_paragraph():
    document = Document()
    paragraph = document.add_paragraph()

    result = ParagraphParser(paragraph).parse()

    assert result == ""


def test_whitespace_is_preserved():
    document = Document()
    paragraph = document.add_paragraph()

    paragraph.add_run("Hello   World")

    result = ParagraphParser(paragraph).parse()

    assert result == "Hello   World"