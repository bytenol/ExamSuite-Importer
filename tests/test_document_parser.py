from pathlib import Path

from docx import Document
from lxml import etree

from examsuite_importer import DocumentParser, QuestionType

def create_docx(tmp_path: Path, paragraphs: list[str]) -> Path:
    file_path = tmp_path / "exam.docx"

    document = Document()

    for text in paragraphs:
        document.add_paragraph(text)

    document.save(file_path)

    return file_path


def test_parse_metadata(tmp_path):
    file_path = create_docx(tmp_path,
        [
        "Subject: Chemistry",
        "Class: SS2",
        "Section: A",
        "Question",
        "Define oxidation.",
        "Mark 2",
        ],
    )

    result = DocumentParser(file_path).parse()

    assert result.success is True
    assert result.metadata.subject == "Chemistry"
    assert result.metadata.class_name == "SS2"
    assert result.metadata.section == "A"


def test_parse_single_theory_question(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Subject: Chemistry",
    "Class: SS2",
    "Section: A",
    "Question",
    "What is oxidation?",
    "Mark 2",
    ],
    )

    result = DocumentParser(file_path).parse()

    assert result.success is True
    assert len(result.questions) == 1

    question = result.questions[0]

    assert question.number == 1
    assert question.content == "What is oxidation?"
    assert question.type == QuestionType.THEORY
    assert question.mark == 2


def test_parse_multiple_questions(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Subject: Chemistry",
    "Class: SS2",
    "Section: A",
    "Question",
    "What is an atom?",
    "Mark 2",
    "Question",
    "Define a molecule.",
    "Mark 3",
    "Question",
    "What is an element?",
    "Mark 1",
    ],
    )


    result = DocumentParser(file_path).parse()

    assert result.success is True
    assert len(result.questions) == 3

    assert result.questions[0].number == 1
    assert result.questions[1].number == 2
    assert result.questions[2].number == 3



def test_question_numbers_are_generated_by_document_parser(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Question",
    "First question",
    "Question",
    "Second question",
    ],
    )


    result = DocumentParser(file_path).parse()

    assert [question.number for question in result.questions] == [1, 2]



def test_blank_paragraphs_are_ignored(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Subject: Chemistry",
    "",
    "",
    "Question",
    "",
    "What is an element?",
    "",
    "Mark 2",
    "",
    ],
    )


    result = DocumentParser(file_path).parse()

    assert result.success is True
    assert len(result.questions) == 1
    assert result.questions[0].content == "What is an element?"


def test_single_choice_question(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Subject: Chemistry",
    "Class: SS1",
    "Section: A",
    "Question",
    "Which of the following is a noble gas?",
    "(a) Oxygen",
    "(b) Nitrogen",
    "(c) Neon",
    "(d) Hydrogen",
    "Answer c",
    "Mark 1",
    ],
    )


    result = DocumentParser(file_path).parse()

    assert result.success is True

    question = result.questions[0]

    assert question.type == QuestionType.SINGLE_CHOICE
    assert question.answer == "c"
    assert question.mark == 1

    assert [
        (choice.label, choice.text)
        for choice in question.choices
    ] == [
        ("a", "Oxygen"),
        ("b", "Nitrogen"),
        ("c", "Neon"),
        ("d", "Hydrogen"),
    ]



def test_fill_in_blank_question(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Question",
    "The chemical formula for water is ______.",
    "Answer H2O",
    "Mark 1",
    ],
    )

    result = DocumentParser(file_path).parse()

    assert result.success is True

    question = result.questions[0]

    assert question.type == QuestionType.FILL_IN_BLANK
    assert question.answer == "H2O"


def test_default_mark_is_used(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Question",
    "Explain diffusion.",
    ],
    )

    result = DocumentParser(file_path).parse()

    assert result.success is True
    assert result.questions[0].mark == 1.0

def test_invalid_question_is_reported(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Question",
    "",
    "Mark 2",
    ],
    )

    
    result = DocumentParser(file_path).parse()

    assert result.success is False
    assert len(result.questions) == 1

    assert any(
        error.question_number == 1
        and error.message == "Question content cannot be empty."
        for error in result.errors
    )
    

def test_invalid_single_choice_question_is_reported(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Question",
    "Which is correct?",
    "(a) Option A",
    "(b) Option B",
    "(c) Option C",
    "Answer a",
    "Mark 1",
    ],
    )


    result = DocumentParser(file_path).parse()

    assert result.success is False

    messages = {
        error.message
        for error in result.errors
    }

    assert (
        "A single-choice question must provide exactly 4 options."
        in messages
    )

    assert (
        "Choices must contain exactly a, b, c, and d."
        in messages
    )
    

def test_multiple_errors_can_belong_to_same_question(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Question",
    "",
    "Mark 0",
    ],
    )

    
    result = DocumentParser(file_path).parse()

    assert result.success is False

    question_errors = [
        error
        for error in result.errors
        if error.question_number == 1
    ]

    assert len(question_errors) >= 2
    

def test_passed_excludes_questions_with_errors(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Question",
    "What is an atom?",
    "Mark 2",
    "Question",
    "",
    "Mark 2",
    "Question",
    "What is a molecule?",
    "Mark 2",
    ],
    )

    
    result = DocumentParser(file_path).parse()

    assert result.success is False
    assert len(result.questions) == 3

    passed_numbers = [
        question.number
        for question in result.passed
    ]

    assert passed_numbers == [1, 3]
    

def test_document_without_questions_returns_success(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "Subject: Chemistry",
    "Class: SS2",
    "Section: A",
    ],
    )


    result = DocumentParser(file_path).parse()

    assert result.success is True
    assert result.questions == []
    assert result.errors == []
    assert result.passed == []
    

def test_text_before_first_question_is_ignored(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "This is an introduction.",
    "Subject: Chemistry",
    "Class: SS2",
    "Question",
    "What is chemistry?",
    ],
    )

    
    result = DocumentParser(file_path).parse()

    assert result.success is True
    assert len(result.questions) == 1
    assert result.questions[0].content == "What is chemistry?"
    

def test_metadata_is_case_insensitive(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "SUBJECT: Chemistry",
    "class: SS2",
    "SeCtIoN: B",
    "Question",
    "Define an acid.",
    ],
    )

    
    result = DocumentParser(file_path).parse()

    assert result.metadata.subject == "Chemistry"
    assert result.metadata.class_name == "SS2"
    assert result.metadata.section == "B"
    

def test_question_marker_is_case_insensitive(tmp_path):
    file_path = create_docx(
    tmp_path,
    [
    "question",
    "What is an element?",
    "QUESTION",
    "What is a compound?",
    ],
    )

    
    result = DocumentParser(file_path).parse()

    assert result.success is True
    assert len(result.questions) == 2
    

def test_formatted_question_content_is_preserved(tmp_path):
    file_path = tmp_path / "formatted.docx"

    
    document = Document()

    document.add_paragraph("Subject: Chemistry")
    document.add_paragraph("Question")

    paragraph = document.add_paragraph()

    run = paragraph.add_run("Define ")
    run.bold = True

    run = paragraph.add_run("oxidation")
    run.italic = True

    run = paragraph.add_run(" and ")
    run.underline = True

    paragraph.add_run("reduction")

    document.save(file_path)

    result = DocumentParser(file_path).parse()

    assert result.success is True

    assert result.questions[0].content == (
        "<strong>Define </strong>"
        "<em>oxidation</em>"
        "<u> and </u>"
        "reduction"
    )
    

def test_equation_is_preserved_as_latex(tmp_path):
    file_path = tmp_path / "equation.docx"

    
    document = Document()

    document.add_paragraph("Subject: Mathematics")
    document.add_paragraph("Question")

    paragraph = document.add_paragraph("Solve ")

    equation = etree.fromstring(
        """
        <m:oMath
            xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
            <m:f>
                <m:fPr/>
                <m:num>
                    <m:r>
                        <m:t>1</m:t>
                    </m:r>
                </m:num>
                <m:den>
                    <m:r>
                        <m:t>2</m:t>
                    </m:r>
                </m:den>
            </m:f>
        </m:oMath>
        """
    )

    paragraph._p.append(equation)

    document.save(file_path)

    result = DocumentParser(file_path).parse()

    assert result.success is True

    content = result.questions[0].content

    assert "1" in content
    assert "2" in content


def test_invalid_file_returns_failed_result(tmp_path):
    file_path = tmp_path / "invalid.docx"

    
    file_path.write_text(
        "This is not a valid DOCX file.",
        encoding="utf-8",
    )

    result = DocumentParser(file_path).parse()

    assert result.success is False
    assert result.questions == []
    assert result.passed == []
    assert len(result.errors) == 1
    assert result.errors[0].question_number is None
    

def test_missing_file_returns_failed_result(tmp_path):
    file_path = tmp_path / "missing.docx"

    
    result = DocumentParser(file_path).parse()

    assert result.success is False
    assert result.questions == []
    assert result.passed == []
    assert len(result.errors) == 1
    assert result.errors[0].question_number is None
    
